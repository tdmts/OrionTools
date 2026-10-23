"""feedback-voorbereiden: de verslagen van een dropbox anoniem, voor Claude."""

import contextlib
import io
import json
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image

from oriontools.export.verslag import kader, voorblad
from oriontools.feedback import voorbereiden

PNG = io.BytesIO()
Image.new("RGB", (120, 80), "red").save(PNG, format="PNG")
PNG = PNG.getvalue()


def verslag(antwoord, figuur=False):
    """Een docx zoals een student het indient: het sjabloon, met een antwoord in het vak."""
    doc = Document()
    voorblad(doc, "Verslag Opdracht RS485", "Je leert RS485.")
    doc.add_paragraph("1. Welke pinnen zijn voor differential signaling?")
    kader(doc, None)
    cel = doc.tables[-1].cell(0, 0)
    cel.paragraphs[0].text = antwoord
    if figuur:
        cel.add_paragraph().add_run().add_picture(io.BytesIO(PNG))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class Basis(unittest.TestCase):
    """Een vak met _feedback/ in .gitignore, en studenten met hun inzendingen."""

    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        (self.root / "oriontools.json").write_text(json.dumps({"course": {"code": "T"}}))
        (self.root / ".gitignore").write_text("_feedback/\n")
        subprocess.run(["git", "init", "-q"], cwd=self.root, check=True)
        self.map = self.root / "_feedback" / "Labo-RS485"
        self.studenten = []

    def student(self, id_, naam, *inzendingen):
        """inzendingen: lijsten van (bestandsnaam, bytes), in volgorde van indienen."""
        subs = []
        for i, bestanden in enumerate(inzendingen):
            files = []
            for naam_bestand, data in bestanden:
                pad = self.map / "bestanden" / naam / f"{i}-{naam_bestand}"
                pad.parent.mkdir(parents=True, exist_ok=True)
                pad.write_bytes(data)
                files.append({"id": i, "name": naam_bestand, "size": len(data), "saved": str(pad)})
            subs.append({"id": i, "date": f"2026-10-0{i + 1}T10:00:00Z", "text": "", "files": files})
        self.studenten.append({"id": id_, "name": naam, "email": None, "submissions": subs, "feedback": None})

    def draai(self):
        self.map.mkdir(parents=True, exist_ok=True)
        (self.map / "inzendingen.json").write_text(json.dumps({
            "orgUnitId": 1, "folderId": 2, "folder": {"name": "Opdracht RS485: Oefeningen (!)"},
            "students": self.studenten}), encoding="utf-8")
        uit = io.StringIO()
        with contextlib.redirect_stdout(uit):
            voorbereiden.main(["_feedback/Labo-RS485", "--repo", str(self.root)])
        return uit.getvalue()

    def md(self, code):
        return (self.map / "anoniem" / code / "verslag.md").read_text(encoding="utf-8")


class FeedbackVoorbereiden(Basis):
    def test_goed_voorblad_weg_antwoord_als_citaat(self):
        self.student(10, "Emma Roose", [("Labo-RS485-verslag.docx", verslag("A en B"))])
        self.draai()
        md = self.md("S01")
        self.assertNotIn("Naam:", md)
        self.assertNotIn("Je leert RS485", md)
        self.assertIn("1. Welke pinnen", md)
        self.assertIn("> A en B", md)

    def test_goed_elke_naam_van_het_vak_wordt_vervangen(self):
        self.student(10, "Emma Roose", [("v.docx", verslag("Samen met Lucas Peeters gedaan, Emma"))])
        self.student(11, "Lucas Peeters", [])
        self.draai()
        md = self.md("S01")
        self.assertIn("Samen met [naam] gedaan, [naam]", md)
        self.assertNotIn("Peeters", md)

    def test_goed_afbeelding_en_code_uit_een_zip(self):
        z = io.BytesIO()
        with zipfile.ZipFile(z, "w") as f:
            f.writestr("map/oef4.ino", "// Emma Roose\nvoid setup() {}\n")
            f.writestr("__MACOSX/map/._oef4.ino", "rommel")
        self.student(10, "Emma Roose", [("v.docx", verslag("Zie figuur", figuur=True))],
                     [("code.zip", z.getvalue())])
        self.draai()
        code = self.map / "anoniem" / "S01" / "code" / "oef4.ino"
        self.assertEqual(code.read_text(encoding="utf-8"), "// [naam]\nvoid setup() {}\n")
        self.assertEqual(len(list((self.map / "anoniem" / "S01" / "code").iterdir())), 1)
        self.assertIn("![afbeelding 1](img/01.png)", self.md("S01"))
        self.assertTrue((self.map / "anoniem" / "S01" / "img" / "01.png").exists())

    def test_goed_getekende_vorm_in_een_vak_blijft_zichtbaar(self):
        from docx.oxml import parse_xml
        doc = Document(io.BytesIO(verslag("")))
        vorm = parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
            '<w:drawing><wp:anchor><wp:extent cx="1" cy="1"/></wp:anchor></w:drawing></w:r>')
        doc.tables[-1].cell(0, 0).paragraphs[0]._p.append(vorm)
        buf = io.BytesIO()
        doc.save(buf)
        self.student(10, "Emma Roose", [("v.docx", buf.getvalue())])
        self.draai()
        self.assertIn("> [vorm]", self.md("S01"))

    def test_goed_figuur_uit_het_sjabloon_wordt_alleen_genoemd(self):
        sjabloon = self.root / "downloads" / "Labo-RS485-verslag.docx"
        sjabloon.parent.mkdir()
        sjabloon.write_bytes(verslag("", figuur=True))
        self.student(10, "Emma Roose", [("v.docx", verslag("Zie figuur", figuur=True))])
        self.draai()
        self.assertIn("[figuur uit het sjabloon]", self.md("S01"))
        self.assertFalse((self.map / "anoniem" / "S01" / "img").exists())

    def test_goed_codes_blijven_bij_een_tweede_run(self):
        self.student(20, "Zoë Wouters", [("v.docx", verslag("x"))])
        self.draai()
        self.student(5, "Adam Claes", [("v.docx", verslag("y"))])
        self.draai()
        sleutel = json.loads((self.map / "sleutel.json").read_text(encoding="utf-8"))["codes"]
        self.assertEqual(sleutel["S01"]["id"], 20)
        self.assertEqual(sleutel["S02"]["id"], 5)

    def test_goed_wie_niets_indiende_krijgt_geen_code(self):
        self.student(10, "Emma Roose", [("v.docx", verslag("x"))])
        self.student(11, "Lucas Peeters")
        self.draai()
        self.assertEqual(list(json.loads((self.map / "sleutel.json").read_text(encoding="utf-8"))["codes"]), ["S01"])

    def test_fout_map_die_git_niet_negeert(self):
        (self.root / ".gitignore").write_text("")
        self.student(10, "Emma Roose", [("v.docx", verslag("x"))])
        with self.assertRaises(SystemExit) as fout:
            self.draai()
        self.assertIn(".gitignore", str(fout.exception))
        self.assertFalse((self.map / "anoniem").exists())

    def test_ontdubbel_laat_kopie_en_pdf_export_weg(self):
        lijst = [("v.docx", b"a"), ("kopie/v.docx", b"a"), ("v.pdf", b"b"), ("ander.pdf", b"c")]
        self.assertEqual(voorbereiden.ontdubbel(lijst), [("v.docx", b"a"), ("ander.pdf", b"c")])


class Spreadsheet(Basis):
    """voorstellen.md naar feedback.xlsx, en de versie van de lector terug naar JSON."""

    VOORSTELLEN = ("# Voorstellen\n\n## S01\n\nGoed\n- Vraag 3: juist.\n\nWerkpunten\n- Vraag 9: waarom.\n\n"
                   "Voor de lector: foto gedraaid.\n\n## S02\n\nGoed\n- Vraag 1: goed.\n\nWerkpunten\n- Vraag 2: fout.\n")

    def setUp(self):
        super().setUp()
        self.student(10, "Zoë Wouters", [("v.docx", verslag("x"))])
        self.student(11, "Adam Claes", [("v.docx", verslag("y"))])
        self.draai()
        (self.map / "voorstellen.md").write_text(self.VOORSTELLEN, encoding="utf-8")

    def commando(self, module, *opties):
        uit = io.StringIO()
        with contextlib.redirect_stdout(uit):
            module.main(["_feedback/Labo-RS485", "--repo", str(self.root), *opties])
        return uit.getvalue()

    def test_goed_rondgang(self):
        from openpyxl import load_workbook
        from oriontools.feedback import naar_json, spreadsheet
        self.commando(spreadsheet)
        wb = load_workbook(self.map / "feedback.xlsx")
        ws = wb["Feedback"]
        rijen = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual([r[0] for r in rijen], ["Adam Claes", "Zoë Wouters"])   # op naam
        zoe = rijen[1]
        self.assertEqual(zoe[1], "S01")
        self.assertNotIn("Voor de lector", zoe[3])
        self.assertEqual(zoe[4], "foto gedraaid.")
        ws.cell(row=3, column=4).value = "Goed\n- Vraag 3: juist, Zoë.\n\nWerkpunten\n- Vraag 9: waarom."
        ws.cell(row=2, column=4).value = None
        wb.save(self.map / "feedback.xlsx")

        self.commando(naar_json)
        data = json.loads((self.map / "feedback.json").read_text(encoding="utf-8"))
        self.assertEqual((data["orgUnitId"], data["folderId"]), (1, 2))
        per = {i["name"]: i for i in data["items"]}
        self.assertEqual(per["Zoë Wouters"]["decision"], "aangepast")
        self.assertEqual(per["Zoë Wouters"]["id"], 10)
        self.assertEqual(per["Adam Claes"]["decision"], "leeg")
        self.assertEqual(per["Adam Claes"]["feedback"], "")
        correcties = (self.map / "correcties.md").read_text(encoding="utf-8")
        self.assertIn("## S01 (aangepast)", correcties)
        self.assertIn("juist, [naam].", correcties)
        self.assertNotIn("Zoë", correcties)
        self.assertNotIn("Adam", correcties)

    def test_fout_bestaande_spreadsheet_niet_overschrijven(self):
        from oriontools.feedback import spreadsheet
        self.commando(spreadsheet)
        with self.assertRaises(SystemExit):
            self.commando(spreadsheet)
        self.commando(spreadsheet, "--opnieuw")

    def test_fout_verschoven_rij(self):
        from openpyxl import load_workbook
        from oriontools.feedback import naar_json, spreadsheet
        self.commando(spreadsheet)
        wb = load_workbook(self.map / "feedback.xlsx")
        wb["Feedback"].cell(row=2, column=1).value = "Zoë Wouters"   # naam bij de verkeerde code
        wb.save(self.map / "feedback.xlsx")
        with self.assertRaises(SystemExit):
            self.commando(naar_json)
        self.assertFalse((self.map / "feedback.json").exists())


if __name__ == "__main__":
    unittest.main()

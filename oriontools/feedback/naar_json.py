"""Lees de spreadsheet van de lector terug, voor npm run dropbox -- feedback.

    python ../OrionTools/orion.py feedback-json _feedback/Labo-RS485

Leest <map>/feedback.xlsx en schrijft twee bestanden:

    <map>/feedback.json     wat OrionSync publiceert:
                            { orgUnitId, folderId, items: [{ id, name, decision, feedback }] }
    <map>/correcties.md     per code het voorstel en de versie van de lector, waar ze
                            verschillen; zonder namen

decision is "voorstel" als de lector de tekst liet staan, "aangepast" als hij
hem veranderde en "leeg" als hij de cel leegmaakte. OrionSync stuurt het niet
mee; het telt ze in zijn proefdraai. Een lege cel slaat OrionSync over. Onder
elke feedback die verstuurd wordt, zet het feedback.vermelding uit de config.

correcties.md is wat Claude leest om de notities bij een opdracht voor te
stellen (stap 7 van de skill): het verschil zonder de namen, die in
feedback.xlsx staan en die Claude niet opent.

Wat de lector veranderde, gaat ook terug in voorstellen.md, met de namen
vervangen zoals in correcties.md. Een volgende ronde bouwt de spreadsheet
opnieuw uit voorstellen.md (feedback-spreadsheet --opnieuw), en zonder die
stap zou de versie van de lector daar verdwijnen en het voorstel opnieuw
klaarstaan. Een cel die de lector leegmaakte, wordt een sectie zonder tekst,
met "De lector stuurde niets." in de regel Voor de lector: de sectie bestaat
dus, zodat ze niet opnieuw geschreven wordt, en ze stuurt niets.

Daarna, in OrionSync, eerst zonder --apply:

    npm run dropbox -- feedback <orgUnitId> <folderId> --in ../<Vak>/<map>/feedback.json
"""

import argparse
import json
import re
import sys

from .. import repo
from .spreadsheet import KOLOMMEN
from .voorbereiden import Namen

NIETS = "De lector stuurde niets."


def gelijk(a, b):
    return " ".join((a or "").split()) == " ".join((b or "").split())


def terugzetten(pad, versies):
    """Zet {code: tekst van de lector} in voorstellen.md; een lege tekst is 'niets sturen'.

    Geeft de codes terug die veranderd zijn. De regel Voor de lector blijft staan.
    """
    if not versies or not pad.is_file():
        return []
    tekst = pad.read_text(encoding="utf-8")
    secties = list(re.finditer(r"^## (S\d+)[ \t]*\n(.*?)(?=^## S\d+[ \t]*\n|\Z)", tekst, re.S | re.M))
    uit, vorig, veranderd = [], 0, []
    for m in secties:
        uit.append(tekst[vorig:m.start()])
        vorig = m.end()
        code = m.group(1)
        if code not in versies:
            uit.append(m.group(0))
            continue
        body = m.group(2).strip()
        noot = ""
        n = re.search(r"^Voor de lector:\s*(.*?)\s*\Z", body, re.S | re.M)
        if n:
            noot, body = n.group(1).strip(), body[:n.start()].strip()
        nieuw = versies[code].replace("\r\n", "\n").strip()
        if not nieuw and not noot.startswith(NIETS):
            noot = f"{NIETS} {noot}".strip()
        delen = [f"## {code}\n"] + ([nieuw + "\n"] if nieuw else []) + ([f"Voor de lector: {noot}\n"] if noot else [])
        nieuwe_sectie = "\n".join(delen) + "\n"
        if gelijk(nieuwe_sectie, m.group(0)):
            uit.append(m.group(0))      # niets veranderd: laat de sectie byte voor byte staan
            continue
        veranderd.append(code)
        uit.append(nieuwe_sectie)
    uit.append(tekst[vorig:])
    if veranderd:
        pad.write_text("".join(uit), encoding="utf-8")
    return veranderd


def main(argv):
    ap = argparse.ArgumentParser(prog="orion.py feedback-json", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("map", help="de map met feedback.xlsx, bv. _feedback/Labo-RS485")
    repo.voeg_repo_toe(ap)
    args = ap.parse_args(argv)
    vak = repo.vind(args)
    map_ = (vak.root / args.map).resolve()

    from openpyxl import load_workbook

    bron = map_ / "feedback.xlsx"
    if not bron.is_file():
        sys.exit(f"Niet gevonden: {bron}. Maak hem met orion.py feedback-spreadsheet.")
    # read_only houdt het bestand open tot close(), en op Windows kan Excel het
    # dan niet meer bewaren.
    wb = load_workbook(bron, read_only=True)
    try:
        rijen = list(wb["Feedback"].iter_rows(values_only=True))
        info = {r[0]: r[1] for r in wb["Dropbox"].iter_rows(values_only=True)}
    finally:
        wb.close()
    if not rijen or list(rijen[0][:len(KOLOMMEN)]) != KOLOMMEN:
        sys.exit(f"De eerste rij van {bron.name} is niet {KOLOMMEN}: is er een kolom verschoven?")

    sleutel = json.loads((map_ / "sleutel.json").read_text(encoding="utf-8"))["codes"]
    namen = Namen(json.loads((map_ / "inzendingen.json").read_text(encoding="utf-8"))["students"])

    vermelding = vak.config["feedback"]["vermelding"].strip()
    items, correcties, fout, versies = [], [], [], {}
    tel = {"voorstel": 0, "aangepast": 0, "leeg": 0}
    for naam, code, voorstel, feedback, *_ in rijen[1:]:
        if not code:
            continue
        wie = sleutel.get(code)
        if wie is None or wie["name"] != naam:
            fout.append(f"{code}: de naam in de rij hoort niet bij die code")
            continue
        feedback = (feedback or "").strip()
        beslissing = "leeg" if not feedback else "voorstel" if gelijk(feedback, voorstel) else "aangepast"
        tel[beslissing] += 1
        verstuurd = f"{feedback}\n\n{vermelding}" if feedback and vermelding else feedback
        items.append({"id": wie["id"], "name": naam, "decision": beslissing, "feedback": verstuurd})
        if beslissing != "voorstel":
            versies[code] = namen(feedback)
            correcties.append(f"## {code} ({beslissing})\n\nVoorstel:\n\n{voorstel}\n\n"
                              f"Lector:\n\n{namen(feedback) or '(niets gestuurd)'}\n")
    if fout:
        sys.exit("Niets geschreven:\n  " + "\n  ".join(fout))

    (map_ / "feedback.json").write_text(json.dumps(
        {"orgUnitId": info["orgUnitId"], "folderId": info["folderId"], "items": items},
        indent=1, ensure_ascii=False), encoding="utf-8")
    (map_ / "correcties.md").write_text(
        f"# Correcties: {info.get('dropbox', '')}\n\n"
        f"{tel['voorstel']} ongewijzigd, {tel['aangepast']} aangepast, {tel['leeg']} leeggemaakt.\n\n"
        + "\n".join(correcties), encoding="utf-8")
    teruggezet = terugzetten(map_ / "voorstellen.md", versies)
    print(f"{len(items)} rijen: {tel['voorstel']} ongewijzigd, {tel['aangepast']} aangepast, "
          f"{tel['leeg']} leeg. Geschreven: feedback.json en correcties.md.")
    if teruggezet:
        print(f"  de versie van de lector staat nu ook in voorstellen.md: {', '.join(teruggezet)}")
    print(f"Proefdraai in OrionSync: npm run dropbox -- feedback {info['orgUnitId']} {info['folderId']} "
          f"--in {map_ / 'feedback.json'}")
    return 0

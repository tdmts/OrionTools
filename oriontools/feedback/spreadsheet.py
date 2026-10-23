"""Zet de voorstellen van een dropbox in een spreadsheet, met de namen terug erbij.

    python ../OrionTools/orion.py feedback-spreadsheet _feedback/Labo-RS485

Leest <map>/voorstellen.md (wat de skill orion-feedback schreef, per code),
<map>/sleutel.json en <map>/inzendingen.json, en schrijft <map>/feedback.xlsx:
een rij per student met een voorstel, gesorteerd op naam, want de lector zoekt
een student op zijn naam en niet op zijn code.

    Naam | Code | Voorstel | Feedback | Voor de lector | Al feedback

Voorstel is wat Claude schreef en blijft staan om mee te vergelijken. Feedback
begint als een kopie en is wat de student krijgt: pas aan, of maak de cel leeg
om niets te sturen. feedback-json leest de spreadsheet terug.

Een spreadsheet die al bestaat, wordt niet overschreven: daar staat werk van de
lector in. --opnieuw overschrijft toch.

Claude opent feedback.xlsx niet: er staan namen in.
"""

import argparse
import json
import re
import sys

from .. import repo

KOLOMMEN = ["Naam", "Code", "Voorstel", "Feedback", "Voor de lector", "Al feedback"]
BREEDTE = {"Naam": 28, "Code": 7, "Voorstel": 70, "Feedback": 70, "Voor de lector": 40, "Al feedback": 12}


def lees_voorstellen(pad):
    """{code: (feedback, noot)} uit voorstellen.md; de noot is de regel 'Voor de lector:'."""
    tekst = pad.read_text(encoding="utf-8")
    uit = {}
    for m in re.finditer(r"^## (S\d+)[ \t]*\n(.*?)(?=^## S\d+[ \t]*\n|\Z)", tekst, re.S | re.M):
        body = m.group(2).strip()
        noot = ""
        n = re.search(r"^Voor de lector:\s*(.*?)\s*\Z", body, re.S | re.M)
        if n:
            noot = n.group(1).strip()
            body = body[:n.start()].strip()
        uit[m.group(1)] = (body, noot)
    return uit


def main(argv):
    ap = argparse.ArgumentParser(prog="orion.py feedback-spreadsheet", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("map", help="de map met voorstellen.md, bv. _feedback/Labo-RS485")
    ap.add_argument("--opnieuw", action="store_true", help="overschrijf een bestaande feedback.xlsx")
    repo.voeg_repo_toe(ap)
    args = ap.parse_args(argv)
    vak = repo.vind(args)
    map_ = (vak.root / args.map).resolve()

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    uit = map_ / "feedback.xlsx"
    if uit.exists() and not args.opnieuw:
        sys.exit(f"{uit.name} bestaat al en kan aanpassingen van de lector bevatten. "
                 f"Draai met --opnieuw om ze te overschrijven.")
    for nodig in ("voorstellen.md", "sleutel.json", "inzendingen.json"):
        if not (map_ / nodig).is_file():
            sys.exit(f"Niet gevonden: {map_ / nodig}")

    voorstellen = lees_voorstellen(map_ / "voorstellen.md")
    sleutel = json.loads((map_ / "sleutel.json").read_text(encoding="utf-8"))
    data = json.loads((map_ / "inzendingen.json").read_text(encoding="utf-8"))
    per_id = {str(s["id"]): s for s in data["students"]}

    rijen, onbekend = [], []
    for code, (feedback, noot) in voorstellen.items():
        wie = sleutel["codes"].get(code)
        if wie is None:
            onbekend.append(code)
            continue
        student = per_id.get(str(wie["id"]), {})
        al = (student.get("feedback") or {}).get("text", "")
        rijen.append([wie["name"], code, feedback, feedback, noot, "ja" if al else ""])
    if onbekend:
        sys.exit(f"Codes in voorstellen.md die sleutel.json niet kent: {', '.join(onbekend)}")
    rijen.sort(key=lambda r: r[0].lower())

    wb = Workbook()
    ws = wb.active
    ws.title = "Feedback"
    ws.append(KOLOMMEN)
    for r in rijen:
        ws.append(r)
    for i, kolom in enumerate(KOLOMMEN, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = BREEDTE[kolom]
        ws.cell(row=1, column=i).font = Font(bold=True)
    for rij in ws.iter_rows(min_row=2):
        for cel in rij:
            cel.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    # De orgUnitId en folderId reizen mee, zodat feedback-json niet naar een
    # andere dropbox kan schrijven dan die waaruit de verslagen kwamen.
    info = wb.create_sheet("Dropbox")
    info.append(["orgUnitId", data["orgUnitId"]])
    info.append(["folderId", data["folderId"]])
    info.append(["dropbox", data["folder"]["name"]])
    wb.save(uit)
    print(f"{len(rijen)} rijen in {uit.relative_to(vak.root)}. Pas de kolom Feedback aan; "
          f"een lege cel stuurt niets. Daarna: orion.py feedback-json {args.map}")
    # feedback-json zet de versie van de lector terug in voorstellen.md, met de
    # namen vervangen; zo'n tekst mag niet ongezien als "[naam]" vertrekken.
    met_naam = sorted(r[1] for r in rijen if "[naam]" in (r[3] or ""))
    if met_naam:
        print(f"  [naam] in de feedback van {', '.join(met_naam)}: vul de naam in of pas de zin aan")
    return 0

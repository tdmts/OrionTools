"""Genereer het verslagsjabloon (.docx) van een labo uit zijn Opdracht.html.

    python ../OrionTools/orion.py export-verslag Labo/RS485

Het pad is de map met de Opdracht.html, relatief tegen de root van de vakrepo.
Het document komt standaard in de downloads-map van het vak (paths.downloads),
onder het pad van die map met streepjes: Labo/RS485 wordt
Labo-RS485-verslag.docx.

Het sjabloon is geen tweede versie van de opdracht die je apart onderhoudt: het
wordt uit de pagina gegenereerd, zodat het niet kan verouderen. Draai dit
commando opnieuw in dezelfde commit als een wijziging aan Opdracht.html, en
commit de docx mee. OrionSync spiegelt alleen bestanden die in git zitten naar
Brightspace.

Waarom een docx en geen invulveld op de pagina: de student moet screenshots
kunnen plakken, moet van pc kunnen wisselen, en er moet een backup van zijn.
Browseropslag doet geen van die drie. Een bestand op zijn OneDrive wel.

Wat het script uit de pagina leest
----------------------------------
De volledige opdracht staat in <!-- verslag ... --> in Opdracht.html en komt
hier terecht. De pagina zelf is een landingspagina: waar het labo over gaat en
de knop om dit document te downloaden. Dat is met opzet. De student werkt in dit
document, en dezelfde opdracht op twee plaatsen laat hem twijfelen waar hij moet
antwoorden. Door ze toch in Opdracht.html te laten staan, kan dit sjabloon niet
afwijken van de opdracht waar het bij hoort, en blijft de opdracht bewerkbaar
met dezelfde HTML-componenten als de rest van de site.

Binnen dat blok:

* h1 / h2 / h3, alinea's, lijsten en codeblokken worden gewone tekst.
* <h2 class="nieuwe-pagina"> begint op een nieuw blad: voor een deel dat de
  student los van de rest afwerkt, zoals de berekeningen die hij thuis maakt.
* <code>, <kbd>, <strong> en <em> binnen een zin houden hun opmaak: een commando
  in Consolas op een grijze achtergrond, een toets vet met een randje. Zonder
  dat was elke alinea platte tekst, en stond "typ show vlan" in het document
  als een zin waarin niets zegt welke woorden je typt.
* .terminal-window wordt een donker vak, zoals op de site: de prompt groen, het
  commando vet en wit, de uitvoer grijs. Zo staat wat de student typt apart van
  wat al op het scherm staat, en toont de prompt van een switch ook in welke
  context hij het typt ((config)#, (vlan-100)#).
* <p class="controle"> wordt een omkaderde oranje balk: het moment waarop de
  student zijn werk laat nakijken. Als vette zin tussen de alinea's las hij er
  overheen.
* <ol class="vragen"> is een vragenlijst: elke <li> wordt een genummerde vraag
  met een leeg kader eronder. Geef een <li> de class "screenshot" en het kader
  wordt groot genoeg om er een in te plakken.
* class="verslag-kader" met data-verslag="..." zet een leeg kader met dat
  opschrift, voor een foto of een screenshot.
* <table class="verslag-tabel"> binnen een <li> wordt een invultabel: de <th>'s
  zijn de kolomkoppen, de eerste <td> van elke rij is het rijlabel, en de rest
  komt leeg. Gebruik dat waar de student per rij een kruisje zet. Zo'n vraag
  krijgt geen los kader meer; de tabel is de invulruimte. Met de class
  "bits" erbij wordt het een vak per bit, zoals de AND in de syllabus: zie
  bittabel().
* <figure> wordt de afbeelding zelf, ingesloten, met haar figcaption eronder.
  Geen verwijzing naar de pagina: daar staat de opdracht niet meer. Een
  <figure> of een <pre> mag ook binnen een <li> staan: het blok krijgt dan
  zijn eigen alinea's onder die stap, in plaats van dat de afbeelding
  wegvalt en haar onderschrift in de zin van de stap belandt. Een png komt
  op ware grootte bij 96 dpi, tot de tekstbreedte; een width op de <img>
  maakt hem kleiner.
* Een .svg wordt onderweg door Chrome tot een png gemaakt. Zie "Een svg in de
  opdracht" hieronder; dat is het enige wat een browser nodig heeft.
* .accordion-item wordt een tip die openstaat. Op het scherm klapt die dicht,
  maar wat hier niet in staat bestaat voor de student niet. Zet een
  .accordion-container binnen de <li> waar de tip over gaat: hij komt dan
  tussen die vraag en haar invulruimte, in plaats van onderaan de lijst waar
  de student hem pas leest als hij het antwoord al opgeschreven heeft.

De nummering loopt door over de secties heen, dus vraag 11 volgt op vraag 10 ook
al staan er twee koppen tussen.

Een svg in de opdracht
----------------------
Word kent geen svg, en python-docx dus ook niet: add_picture() weigert er een.
Daarom wordt een svg hier eerst door een headless Chrome gehaald, precies zoals
de syllabus en de handout dat doen, en gaat de png die daaruit komt het document
in.

Die png wordt NIET bewaard en niet gecommit. Hij ontstaat in een tijdelijke map
en verdwijnt zodra dit script klaar is. Dat is de hele reden dat het zo werkt:
de svg blijft de enige bron. Een png ernaast bewaren zou betekenen dat een
figuur twee bestanden heeft die uit elkaar kunnen lopen, en dat is precies wat
het hertekenen van img/ probeert weg te werken.

De raster wordt op drie keer de eigen maat getrokken, want in de docx is het
weer gewoon een raster en de student drukt dat document af. De breedte in het
document komt niet uit die pixels maar uit de maat die de svg zelf opgeeft, op
96 dpi, net als bij een png. Anders zou de schaalfactor de figuur breed maken.

Chrome is dus alleen nodig als er echt een svg in de opdracht staat. Een labo
zonder svg draait met niets dan python-docx, en dat blijft zo.

De regels unclosed-comment en verslag-markup van de check (orion.py check)
bewaken dat het blok afgesloten is en dat er geen vragen buiten terechtkomen;
verslag-controle dat elk controlemoment een <p class="controle"> is.

Nodig: python-docx (pip install python-docx), en een headless Chrome of Edge
zodra de opdracht een svg bevat. Waar die staat weet oriontools/chrome.py.
"""

import argparse
import html
import re
import subprocess
import sys
import tempfile
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlparse, unquote

from .. import repo
from ..chrome import zoek_chrome

try:
    from docx import Document
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:
    sys.exit("python-docx ontbreekt. Installeer het met: pip install python-docx")


# --------------------------------------------------------------------------
# HTML -> boom
# --------------------------------------------------------------------------

VOID = {"img", "br", "hr", "input", "meta", "link"}

# Het opschrift waarmee een HTML-commentaar zegt dat het in het verslag hoort.
VERSLAG_MARKER = "verslag"
SKIP_CONTENT = {"script", "style", "head"}


class Node:
    def __init__(self, tag, attrs=None):
        self.tag = tag
        self.attrs = dict(attrs or {})
        self.children = []
        self.text = ""

    @property
    def classes(self):
        return set((self.attrs.get("class") or "").split())

    def find_all(self, tag):
        for child in self.children:
            if child.tag == tag:
                yield child
            yield from child.find_all(tag)

    def _flat_text(self):
        """Alle tekst eronder, aan elkaar, met de witruimte nog intact.

        Hier mag niet gestript worden: de spatie tussen "Zoek de" en
        "<em>pin configuration</em>" zit in het tekstknoopje ervoor, en wie die
        per knoop wegstript plakt de woorden aan elkaar.
        """
        parts = [self.text] if self.text else []
        for child in self.children:
            if child.tag in SKIP_CONTENT:
                continue
            parts.append(child._flat_text())
        return "".join(parts)

    def inner_text(self):
        """Platte tekst, met de spaties genormaliseerd zoals een browser dat doet."""
        return re.sub(r"\s+", " ", self._flat_text()).strip()

    def raw_text(self):
        """Tekst met de witruimte intact, voor codeblokken."""
        parts = [self.text] if self.text else []
        for child in self.children:
            parts.append(child.raw_text())
        return "".join(parts)


class TreeBuilder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = Node(tag, attrs)
        self.stack[-1].children.append(node)
        if tag not in VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.stack[-1].children.append(Node(tag, attrs))

    def handle_endtag(self, tag):
        if tag in VOID:
            return
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        node = Node("#text")
        node.text = data
        self.stack[-1].children.append(node)

    def handle_comment(self, data):
        # Op zijn plaats in de boom, want de volgorde van het verslag is de
        # volgorde van de pagina.
        node = Node("#comment")
        node.text = data
        self.stack[-1].children.append(node)


def parse_fragment(markup):
    builder = TreeBuilder()
    builder.feed(markup)
    return builder.root


def parse(path):
    return parse_fragment(path.read_text(encoding="utf-8"))


# --------------------------------------------------------------------------
# docx
# --------------------------------------------------------------------------

GRIJS = RGBColor(0x5A, 0x5A, 0x5A)

# Wat een tag binnen een zin betekent. De andere tags dragen geen opmaak en
# leveren alleen hun tekst.
INLINE = {"code": "code", "kbd": "kbd", "strong": "vet", "b": "vet",
          "em": "cursief", "i": "cursief"}


def segmenten(node):
    """De tekst van een knoop als (tekst, stijlen)-paren.

    De witruimte wordt samengevoegd zoals inner_text() dat doet, ook over de
    grens van twee tags heen: "typ <code>ls</code> en" houdt een spatie aan
    elke kant van ls, en niet twee.
    """
    ruw = []

    def loop(n, stijl):
        for c in n.children:
            if c.tag in SKIP_CONTENT or c.tag == "#comment":
                continue
            if c.tag == "#text":
                ruw.append((c.text, stijl))
            else:
                extra = INLINE.get(c.tag)
                loop(c, stijl | {extra} if extra else stijl)

    loop(node, frozenset())
    uit = []
    spatie = True
    for tekst, stijl in ruw:
        tekst = re.sub(r"\s+", " ", tekst)
        if spatie:
            tekst = tekst.lstrip(" ")
        if tekst:
            uit.append([tekst, stijl])
            spatie = tekst.endswith(" ")
    while uit:
        uit[-1][0] = uit[-1][0].rstrip(" ")
        if uit[-1][0]:
            break
        uit.pop()
    return [(tekst, stijl) for tekst, stijl in uit]


def vul_achter(pr, kleur):
    """Een achtergrondkleur op een run, een alinea of een cel (w:shd)."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), kleur)
    pr.append(shd)


def inline(p, node, vet=False):
    """De tekst van node als runs in alinea p, met de opmaak van INLINE."""
    for tekst, stijl in segmenten(node):
        run = p.add_run(tekst)
        if vet or "vet" in stijl or "kbd" in stijl:
            run.bold = True
        if "cursief" in stijl:
            run.italic = True
        if "code" in stijl:
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        if "kbd" in stijl:
            run.font.size = Pt(9)
        # De volgorde in w:rPr ligt vast (bdr voor shd, allebei na sz), dus
        # die twee komen er pas bij als python-docx klaar is met de rest.
        if "kbd" in stijl:
            rand = OxmlElement("w:bdr")
            rand.set(qn("w:val"), "single")
            rand.set(qn("w:sz"), "4")
            rand.set(qn("w:space"), "0")
            rand.set(qn("w:color"), "000000")
            run._r.get_or_add_rPr().append(rand)
        if "code" in stijl:
            vul_achter(run._r.get_or_add_rPr(), "EDEDED")


TERMINAL = "1E1E1E"
PROMPT = RGBColor(0x75, 0xA3, 0x57)
COMMANDO = RGBColor(0xFF, 0xFF, 0xFF)
UITVOER = RGBColor(0xCC, 0xCC, 0xCC)


def terminal(doc, node):
    """Een .terminal-window als een donkere cel, met de kleuren van OrionCSS."""
    tabel = doc.add_table(rows=1, cols=1)
    # Niet over een bladgrens: het commando onderaan de ene bladzijde en het
    # antwoord van het toestel bovenaan de volgende lezen als twee dingen.
    tabel.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    cel = tabel.rows[0].cells[0]
    vul_achter(cel._tc.get_or_add_tcPr(), TERMINAL)
    alineas = []

    def regel():
        p = cel.paragraphs[0] if not alineas else cel.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        alineas.append(p)
        return p

    def mono(run, kleur, vet=False):
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = kleur
        run.bold = vet

    for kind in node.children:
        if "term-line" in kind.classes:
            p = regel()
            for span in kind.children:
                if "term-prompt" in span.classes:
                    mono(p.add_run(span.inner_text() + " "), PROMPT, vet=True)
                elif "term-cmd" in span.classes:
                    mono(p.add_run(span.inner_text()), COMMANDO, vet=True)
        elif "term-out" in kind.classes:
            for lijn in html.unescape(kind.raw_text()).strip("\n").split("\n"):
                mono(regel().add_run(lijn if lijn.strip() else " "), UITVOER)
    if alineas:
        alineas[0].paragraph_format.space_before = Pt(4)
        alineas[-1].paragraph_format.space_after = Pt(4)
    doc.add_paragraph()


CONTROLE_RAND = "E36C0A"
CONTROLE_VULLING = "FDEADA"


def controle(doc, node):
    """<p class="controle">: een omkaderde balk over de volle breedte."""
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    randen = OxmlElement("w:pBdr")
    for kant in ("top", "left", "bottom", "right"):
        rand = OxmlElement(f"w:{kant}")
        rand.set(qn("w:val"), "single")
        rand.set(qn("w:sz"), "12")
        rand.set(qn("w:space"), "4")
        rand.set(qn("w:color"), CONTROLE_RAND)
        randen.append(rand)
    # pBdr en shd staan in w:pPr voor spacing; python-docx zet spacing daarna
    # zelf op de juiste plaats.
    ppr.append(randen)
    vul_achter(ppr, CONTROLE_VULLING)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(14)
    inline(p, node, vet=True)


def kader(doc, opschrift, hoogte_cm=4.0):
    """Een leeg, omrand vak om in te typen of een screenshot in te plakken."""
    if opschrift:
        label = doc.add_paragraph()
        run = label.add_run(opschrift)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = GRIJS
        label.paragraph_format.space_after = Pt(2)

    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    # De hoogte hoort op de rij, niet op de cel: _Cell heeft geen height, dus
    # cell.height = ... zet een Python-attribuut dat nooit in het document
    # belandt. AT_LEAST laat het vak meegroeien met wat de student erin plakt.
    rij = table.rows[0]
    rij.height = Cm(hoogte_cm)
    rij.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    doc.add_paragraph()


PNG_SIGNATUUR = bytes([0x89, 0x50, 0x4E, 0x47, 0x0D, 0x0A, 0x1A, 0x0A])


def png_afmeting(pad):
    """Breedte en hoogte in pixels uit de PNG-header, zonder Pillow erbij te halen.

    De IHDR-chunk staat altijd vooraan: 8 bytes signatuur, 4 lengte, 4 type,
    dan breedte en hoogte als big-endian 32-bit.
    """
    with open(pad, "rb") as f:
        kop = f.read(24)
    if len(kop) < 24 or kop[:8] != PNG_SIGNATUUR:
        return None
    return (int.from_bytes(kop[16:20], "big"), int.from_bytes(kop[20:24], "big"))


# De tekstbreedte van een A4 met de standaardmarges van Word.
MAX_BREEDTE_CM = 15.0

# Hoeveel keer zijn eigen maat de svg getrokken wordt. In het document is het
# weer een raster, en de student drukt het af.
SVG_SCHAAL = 3

SVG_TAG_RE = re.compile(r"<svg\b[^>]*>", re.I | re.S)
# Niet \b voor de naam: in stroke-width staat ook een woordgrens, en dan
# wordt de lijndikte de breedte van de figuur.
MAAT_RE = re.compile(r'(?<![-\w])(width|height)\s*=\s*["\']([0-9.]+)', re.I)
VIEWBOX_RE = re.compile(r'viewBox\s*=\s*["\']\s*[-0-9.]+\s+[-0-9.]+\s+'
                        r'([0-9.]+)\s+([0-9.]+)', re.I)


def svg_maat(pad):
    """Breedte en hoogte van een svg in CSS-pixels.

    De eigen reeks in img/ zet width en height allebei op de <svg>. Staan ze
    er niet, dan geeft de viewBox dezelfde twee getallen. Zonder een van
    beide valt er niets te tekenen, want dan weet ook de browser de maat niet.
    """
    tag = SVG_TAG_RE.search(pad.read_text(encoding="utf-8", errors="replace"))
    if not tag:
        sys.exit(f"{pad.name}: geen <svg>-tag gevonden")
    # Alleen binnen de openingstag zoeken. Elke <rect> eronder draagt ook
    # een width, en die van de eerste zou anders de maat bepalen.
    tag = tag.group(0)
    maten = {}
    for m in MAAT_RE.finditer(tag):
        maten.setdefault(m.group(1).lower(), float(m.group(2)))
    if "width" in maten and "height" in maten:
        return maten["width"], maten["height"]
    vak = VIEWBOX_RE.search(tag)
    if vak:
        return float(vak.group(1)), float(vak.group(2))
    sys.exit(f"{pad.name}: geen width/height en geen viewBox, dus geen maat")


def svg_naar_png(pad, doel):
    """Trek de svg met Chrome tot een png op SVG_SCHAAL keer zijn eigen maat."""
    breedte, hoogte = svg_maat(pad)
    opdracht = [zoek_chrome(), "--headless=new", "--disable-gpu",
                "--no-sandbox", "--hide-scrollbars",
                f"--force-device-scale-factor={SVG_SCHAAL}",
                f"--window-size={int(round(breedte))},{int(round(hoogte))}",
                "--default-background-color=FFFFFFFF",
                f"--screenshot={Path(doel).resolve()}",
                pad.resolve().as_uri()]
    klaar = subprocess.run(opdracht, capture_output=True, text=True)
    if not Path(doel).exists():
        sys.exit(f"Chrome maakte geen png van {pad.name}:\n{klaar.stderr[-2000:]}")
    return breedte


def figuur(doc, node, basis):
    """Sluit de afbeelding in. Niet als verwijzing: dit document is de opdracht.

    Zolang de opdracht ook op de pagina stond, volstond een regel "zie de
    figuur op de opdrachtpagina". Nu staat ze daar niet meer, dus een
    verwijzing zou naar niets wijzen.
    """
    imgs = list(node.find_all("img"))
    if not imgs:
        return
    src = imgs[0].attrs.get("src", "")
    alt = imgs[0].attrs.get("alt", "")
    pad = (basis / unquote(urlparse(src).path)).resolve()

    if not pad.exists():
        p = doc.add_paragraph()
        run = p.add_run(f"[Afbeelding ontbreekt: {src}]")
        run.italic = True
        run.font.color.rgb = GRIJS
        return

    if pad.suffix.lower() == ".svg":
        # De png bestaat alleen tijdens deze aanroep; de svg blijft de bron.
        with tempfile.TemporaryDirectory() as tijdelijk:
            png = Path(tijdelijk) / "figuur.png"
            breedte_px = svg_naar_png(pad, png)
            # Uit de maat van de svg, niet uit de pixels: die dragen SVG_SCHAAL.
            breedte = Cm(min(breedte_px / 96 * 2.54, MAX_BREEDTE_CM))
            doc.add_picture(str(png), width=breedte)
    else:
        breedte = Cm(MAX_BREEDTE_CM)
        afmeting = png_afmeting(pad)
        if afmeting:
            # 96 dpi is wat Word aanneemt. Een klein plaatje niet opblazen.
            natuurlijk_cm = afmeting[0] / 96 * 2.54
            breedte = Cm(min(natuurlijk_cm, MAX_BREEDTE_CM))
        # Een width op de <img> verkleint een screenshot die op ware grootte
        # een halve bladzijde vult, zoals een Windows-dialoogvenster. Ook in
        # CSS-pixels, dus 96 dpi; de pixels zelf blijven, zodat de tekst
        # scherp blijft.
        opgegeven = imgs[0].attrs.get("width")
        if opgegeven:
            breedte = Cm(min(float(opgegeven) / 96 * 2.54, MAX_BREEDTE_CM))
        doc.add_picture(str(pad), width=breedte)

    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    onderschriften = [c for c in node.children if c.tag == "figcaption"]
    tekst = onderschriften[0].inner_text() if onderschriften else alt
    if tekst:
        p = doc.add_paragraph()
        run = p.add_run(tekst)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = GRIJS
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def accordion(doc, node):
    """Een hint of een tip. Op het scherm klapt die open; hier staat hij open.

    Weglaten kan niet meer: de opdracht staat alleen nog in dit document, dus
    wat hier niet in staat, bestaat voor de student niet.
    """
    for item in node.children:
        if "accordion-item" not in item.classes:
            continue
        titels = [c for c in item.children if "title" in c.classes]
        rest = Node("div")
        rest.children = [c for c in item.children if "title" not in c.classes]
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        if titels:
            p.add_run(titels[0].inner_text() + " ").bold = True
        inline(p, rest)


def codeblok(doc, code):
    for regel in code.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Cm(0.6)
        run = p.add_run(regel if regel.strip() else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
    doc.add_paragraph()


def nieuwe_nummering(doc, stijl):
    """Een numId die bij 1 begint, op de nummering van de stijl.

    Alle alinea's in "List Number" delen in Word één nummering, dus een tweede
    stappenlijst telde door waar de eerste stopte: de stappen van Scannen in
    Draadloos netwerk begonnen bij 6, na de vijf van de reset (8 oktober 2026).
    Een eigen w:num met een startOverride op dezelfde abstracte nummering
    begint opnieuw bij 1 en houdt de opmaak van de stijl.
    """
    numpr = doc.styles[stijl].element.find(qn("w:pPr") + "/" + qn("w:numPr"))
    if numpr is None:
        return None
    numid = numpr.find(qn("w:numId")).get(qn("w:val"))
    nummering = doc.part.numbering_part.element
    num = next(n for n in nummering.findall(qn("w:num")) if n.get(qn("w:numId")) == numid)
    abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    nieuw_id = str(max(int(n.get(qn("w:numId"))) for n in nummering.findall(qn("w:num"))) + 1)
    nieuw = OxmlElement("w:num")
    nieuw.set(qn("w:numId"), nieuw_id)
    a = OxmlElement("w:abstractNumId")
    a.set(qn("w:val"), abstract)
    nieuw.append(a)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    nieuw.append(override)
    nummering.append(nieuw)
    return nieuw_id


def nummer(alinea, numid):
    """Hangt een alinea aan een eigen nummering (zie nieuwe_nummering)."""
    ppr = alinea._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), numid)
    numpr.append(ilvl)
    numpr.append(nid)
    ppr.append(numpr)


def lijst(doc, node, stijl, basis):
    numid = nieuwe_nummering(doc, stijl) if stijl.startswith("List Number") else None
    for li in node.children:
        if li.tag != "li":
            continue
        # Een geneste lijst, figuur of codeblok wordt zelf gerenderd, niet mee
        # in de tekst van de li. Zonder dit verdwijnt de afbeelding en komt haar
        # onderschrift midden in de zin van de stap terecht, zonder dat iets
        # faalt. Een stappenlijst met een screenshot per stap is precies wat een
        # opdracht met een klikpad nodig heeft.
        genest = [c for c in li.children if is_blok(c)]
        eigen = Node("li")
        eigen.children = [c for c in li.children if not is_blok(c)]
        if segmenten(eigen):
            alinea = doc.add_paragraph(style=stijl)
            if numid:
                nummer(alinea, numid)
            inline(alinea, eigen)
        for sub in genest:
            blok_in_li(doc, sub, basis)


BLOKKEN_IN_LI = ("ol", "ul", "figure", "pre")


def is_blok(node):
    """Krijgt dit kind van een <li> zijn eigen alinea's? Een terminalvenster is
    een <div>, dus voor hem hangt het aan de klasse en niet aan de tag."""
    return node.tag in BLOKKEN_IN_LI or "terminal-window" in node.classes


def blok_in_li(doc, sub, basis):
    """Een blok dat binnen een <li> staat en zijn eigen alinea's krijgt."""
    if "terminal-window" in sub.classes:
        terminal(doc, sub)
    elif sub.tag == "figure":
        figuur(doc, sub, basis)
    elif sub.tag == "pre":
        codeblok(doc, html.unescape(sub.raw_text()))
    else:
        lijst(doc, sub, "List Number 2" if sub.tag == "ol" else "List Bullet 2", basis)


def vragenlijst(doc, node, teller, basis):
    for li in node.children:
        if li.tag != "li":
            continue
        teller[0] += 1

        # Een tabel in de vraag is de invulruimte, en hoort dus niet in de
        # vraagtekst zelf. Zelfde aanpak als bij een geneste lijst. Een hint
        # hoort er evenmin in: die staat als aparte alinea onder de vraag.
        tabellen = [c for c in li.children if c.tag == "table"]
        hints = [c for c in li.children if "accordion-container" in c.classes]
        blokken = [c for c in li.children if is_blok(c)]
        eigen = Node("li")
        eigen.children = [c for c in li.children
                          if c.tag != "table" and not is_blok(c)
                          and "accordion-container" not in c.classes]

        p = doc.add_paragraph()
        run = p.add_run(f"{teller[0]}. ")
        run.bold = True
        inline(p, eigen)

        # Een figuur of een codeblok hoort bij de vraag en komt er meteen
        # onder, voor de hint en voor de invulruimte.
        for blok in blokken:
            blok_in_li(doc, blok, basis)

        # De hint komt voor de invulruimte: de student leest hem terwijl hij
        # nog aan het antwoord werkt, niet nadat hij het opgeschreven heeft.
        for hint in hints:
            accordion(doc, hint)

        if tabellen:
            for tabel in tabellen:
                verslag_tabel(doc, tabel)
            continue

        # Een vraag die om een screenshot vraagt, krijgt een vak waar er ook
        # een in past. Twee regels tekst en een screenshot zijn niet dezelfde
        # hoeveelheid plaats.
        screenshot = "screenshot" in li.classes
        kader(doc, "Plak hier je screenshot" if screenshot else None,
              hoogte_cm=6.5 if screenshot else 2.5)


def verslag_tabel(doc, node):
    """Een invultabel: kolomkoppen uit de thead, rijlabels uit de eerste kolom.

    Bedoeld voor een vraag waar de student per rij een kruisje zet. De cellen
    achter het label komen leeg in het document; de tabel zelf is dan de
    invulruimte, en er hoort geen los kader meer bij.
    """
    koppen = [c.inner_text() for c in node.find_all("th")]
    rijen = [r for r in node.find_all("tr") if any(c.tag == "td" for c in r.children)]
    if not koppen or not rijen:
        return

    table = doc.add_table(rows=len(rijen) + 1, cols=len(koppen))
    table.style = "Table Grid"

    for kolom, kop in enumerate(koppen):
        cel = table.cell(0, kolom)
        cel.text = ""
        cel.paragraphs[0].add_run(kop).bold = True

    for n, rij in enumerate(rijen, start=1):
        cellen = [c for c in rij.children if c.tag == "td"]
        for kolom in range(len(koppen)):
            tekst = cellen[kolom].inner_text() if kolom < len(cellen) else ""
            cel = table.cell(n, kolom)
            cel.text = ""
            if tekst:
                cel.paragraphs[0].add_run(tekst)
        table.rows[n].height = Cm(0.8)
        table.rows[n].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    if "bits" in node.classes:
        bittabel(table)

    doc.add_paragraph()


# Breed genoeg voor "Subnetmasker" of "Default gateway"; de rest van de
# tekstbreedte gaat naar de bits.
BITS_LABEL_CM = 2.8


def bittabel(table):
    """Een verslag-tabel met class "bits": een label en dan een vak per bit.

    Even brede kolommen met de marges van Word laten in een vak van 0,45 cm
    geen plaats voor een cijfer, en het label past niet. Dus krijgt het label
    zijn breedte, delen de bits de rest, valt de marge in een cel weg en staat
    er na elke acht bits een dikkere lijn, waar in de decimale schrijfwijze
    het punt staat. De kop is klein gezet: daar staat het gewicht van de bit,
    128 tot 1, en dat is drie cijfers in een vak voor één.
    """
    kolommen = len(table.columns)
    bit_cm = (MAX_BREEDTE_CM - BITS_LABEL_CM) / (kolommen - 1)
    table.autofit = False

    tblpr = table._tbl.tblPr
    marges = OxmlElement("w:tblCellMar")
    for kant in ("left", "right"):
        m = OxmlElement(f"w:{kant}")
        m.set(qn("w:w"), "0")
        m.set(qn("w:type"), "dxa")
        marges.append(m)
    # Het schema legt de volgorde in tblPr vast, en Word negeert een
    # tblCellMar die achter tblLook staat.
    look = tblpr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(marges)
    else:
        tblpr.append(marges)

    for kolom in range(kolommen):
        breedte = Cm(BITS_LABEL_CM if kolom == 0 else bit_cm)
        table.columns[kolom].width = breedte
        for n, rij in enumerate(table.rows):
            cel = rij.cells[kolom]
            cel.width = breedte
            alinea = cel.paragraphs[0]
            if kolom == 0:
                alinea.paragraph_format.left_indent = Cm(0.1)
                continue
            alinea.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if n == 0:
                for run in alinea.runs:
                    run.font.size = Pt(6)
                    run.bold = False
            if kolom % 8 == 1:
                randen = OxmlElement("w:tcBorders")
                links = OxmlElement("w:left")
                links.set(qn("w:val"), "single")
                links.set(qn("w:sz"), "12")
                links.set(qn("w:color"), "000000")
                randen.append(links)
                cel._tc.get_or_add_tcPr().append(randen)


def render(container, doc, basis, teller=None):
    # Eén teller voor het hele document, ook door de commentaarblokken heen,
    # anders begint de nummering opnieuw bij elke sectie.
    if teller is None:
        teller = [0]

    for node in container.children:
        klassen = node.classes

        if node.tag in SKIP_CONTENT:
            continue

        if "accordion-container" in klassen:
            accordion(doc, node)
            continue

        # De vragen en de invulkaders staan in <!-- verslag ... --> op de
        # pagina: ze horen in dit document thuis en niet op het scherm, waar
        # een student ze zou lezen zonder ergens te kunnen antwoorden. Ze
        # staan wel in Opdracht.html, op hun plaats, zodat dit sjabloon niet
        # kan afwijken van de opdracht waar het bij hoort.
        if node.tag == "#comment":
            inhoud = node.text.strip()
            if inhoud.startswith(VERSLAG_MARKER):
                render(parse_fragment(inhoud[len(VERSLAG_MARKER):]), doc, basis, teller)
            continue

        if "verslag-kader" in klassen:
            kader(doc, node.attrs.get("data-verslag") or "Plak hier je screenshot", 6.5)
            continue

        if node.tag == "h1":
            doc.add_heading(node.inner_text(), level=0)
        elif node.tag == "h2":
            kop = doc.add_heading(node.inner_text(), level=1)
            if "nieuwe-pagina" in klassen:
                kop.paragraph_format.page_break_before = True
        elif node.tag == "h3":
            doc.add_heading(node.inner_text(), level=2)
        elif node.tag == "p":
            if "controle" in klassen:
                controle(doc, node)
            elif segmenten(node):
                inline(doc.add_paragraph(), node)
        elif node.tag == "pre":
            codeblok(doc, html.unescape(node.raw_text()))
        elif node.tag == "ol":
            if "vragen" in klassen:
                vragenlijst(doc, node, teller, basis)
            else:
                lijst(doc, node, "List Number", basis)
        elif node.tag == "ul":
            lijst(doc, node, "List Bullet", basis)
        elif node.tag == "figure":
            figuur(doc, node, basis)
        elif node.tag in ("div", "section"):
            if "terminal-window" in klassen:
                terminal(doc, node)
            elif "info-box" in klassen:
                titels = [c for c in node.children if "info-title" in c.classes]
                opschrift = titels[0].inner_text() if titels else ""
                rest = Node("div")
                rest.children = [c for c in node.children if "info-title" not in c.classes]
                p = doc.add_paragraph()
                if opschrift:
                    # Een punt erbij als het opschrift er zelf geen heeft: het
                    # is vet, maar het loopt in dezelfde alinea door in de
                    # tekst erna, en twee zinnen zonder scheiding lezen als een.
                    if opschrift[-1] not in ".:?!":
                        opschrift += "."
                    p.add_run(opschrift + " ").bold = True
                inline(p, rest)
                p.paragraph_format.left_indent = Cm(0.5)
            elif "download-container" in klassen:
                continue
            else:
                render(node, doc, basis, teller)

    return doc


def voorblad(doc, titel, intro):
    """De eerste bladzijde: titel, invulregel, en de introductie van het labo.

    De introductie hoort hier omdat dit document voor de student de opdracht
    is. Wie het opent zonder de site gezien te hebben, moet meteen weten waar
    het labo over gaat. De regel opdracht-lead van de check eist daarom dat
    Opdracht.html een <p class="lead"> heeft.
    """
    doc.add_heading(titel, level=0)

    p = doc.add_paragraph()
    p.add_run("Naam: ").bold = True
    p.add_run(" " * 40)
    p.add_run("     Groep: ").bold = True
    p.add_run(" " * 15)
    p.add_run("     Datum: ").bold = True

    if intro:
        doc.add_paragraph(intro)

    # Alleen wat de student van ons moet horen: invullen en indienen. Waarom dit
    # een document is en geen invulveld op een pagina, is een afweging van ons
    # (patroon 18), en waar hij het bewaart is zijn zaak (18 ook, het voorschrift).
    p = doc.add_paragraph()
    run = p.add_run(
        "Vul dit document in terwijl je werkt. "
        "Je dient het in via de opdracht op Orion."
    )
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GRIJS

    doc.add_page_break()


def main(argv):
    ap = argparse.ArgumentParser(prog="orion.py export-verslag", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("module", help="pad naar de module, bv. Labo/RS485")
    ap.add_argument("-o", "--output", help="pad van de docx, relatief tegen de vakrepo "
                                           "(standaard <downloads>/<Module>-verslag.docx)")
    repo.voeg_repo_toe(ap)
    args = ap.parse_args(argv)
    vak = repo.vind(args)
    wortel = vak.root

    module = (wortel / args.module).resolve()
    opdracht = module / "Opdracht.html"
    if not opdracht.exists():
        sys.exit(f"Niet gevonden: {opdracht}")

    root = parse(opdracht)
    containers = [n for n in root.find_all("div") if "container" in n.classes]
    if not containers:
        sys.exit(f"Geen <div class=\"container\"> in {opdracht}")
    container = containers[0]

    titels = list(container.find_all("h1"))
    titel = titels[0].inner_text() if titels else module.name

    leads = [n for n in container.find_all("p") if "lead" in n.classes]
    if not leads:
        sys.exit(f'{opdracht.name} heeft geen <p class="lead">, dus het document '
                 f"zou beginnen zonder te zeggen waar het labo over gaat")
    intro = leads[0].inner_text()

    # Het volledige pad van de module, niet alleen de laatste twee stukken:
    # Labo/ManagedSwitch/ProCurve en Labo/ManagedSwitch/PacketTracer zouden
    # anders allebei op ManagedSwitch-...-verslag.docx uitkomen.
    if args.output:
        uit = wortel / args.output
    else:
        uit = vak.pad("downloads") / ("-".join(module.relative_to(wortel).parts)
                                      + "-verslag.docx")
    uit.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    voorblad(doc, f"Verslag {titel}", intro)

    # Het document is de <h1>, de lead en het <!-- verslag --> blok. Verder
    # niets. Alles wat daarbuiten op de landingspagina staat is schermwerk
    # (de downloadknop, een verwijzing naar de theorie), en dat hoort niet in
    # een document dat de student offline invult: het verwijst naar een pagina
    # die hij op dat moment niet open heeft, en het vertelt hem iets over hoe
    # het materiaal georganiseerd is in plaats van over zijn labo.
    body = Node("div")
    body.children = [c for c in container.children
                     if c.tag == "#comment" and c.text.strip().startswith(VERSLAG_MARKER)]
    render(body, doc, opdracht.parent)

    try:
        doc.save(uit)
    except PermissionError:
        # Word houdt een geopend document vergrendeld. Dat is geen fout in dit
        # script, en een traceback zegt niet wat je eraan doet.
        sys.exit(f"Kan {uit.name} niet overschrijven. Staat het open in Word? "
                 f"Sluit het en draai dit opnieuw.")

    try:
        toon = uit.relative_to(wortel)
    except ValueError:
        toon = uit          # -o kan naar buiten de repo wijzen
    print(f"Geschreven: {toon}")
    return 0
